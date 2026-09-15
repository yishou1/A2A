import json
import tempfile
import unittest
from pathlib import Path

from src.synapserag.document_ingestion import chunk_document, parse_document
from src.synapserag.document_ingestion.parsers import DocumentParserError


class DocumentIngestionTests(unittest.TestCase):
    def test_markdown_heading_and_deterministic_chunks(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "说明.md"
            path.write_text("# 总则\n\n这是第一段。\n\n## 范围\n\n这是第二段。", encoding="utf-8")
            document = parse_document(path)
            chunks = chunk_document(document, max_tokens=20, overlap_tokens=2)

            self.assertEqual(document.parser, "markdown")
            self.assertEqual(chunks[0].heading_path, ["总则"])
            self.assertEqual(chunks[1].heading_path, ["总则", "范围"])
            self.assertEqual(chunks[0].chunk_id, f"{document.document_id}-chunk-00000")
            self.assertNotEqual(chunks[0].content_sha256, document.content_sha256)
            self.assertIn("总则", chunks[0].text)
            self.assertGreaterEqual(chunks[0].source_segments[0]["source_start"], 0)
            self.assertGreater(chunks[0].source_segments[0]["source_end"], 0)

    def test_text_encoding_and_artifact_shape(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "材料.txt"
            path.write_bytes("第一行\n\n第二行".encode("gb18030"))
            document = parse_document(path, document_id="fixed-id")
            self.assertEqual(document.document_id, "fixed-id")
            self.assertIn("第二行", document.text)
            self.assertEqual(json.loads(json.dumps(document.to_dict(), ensure_ascii=False))["document_id"], "fixed-id")

    def test_unsupported_extension_is_explicit(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "data.xlsx"
            path.write_bytes(b"not supported")
            with self.assertRaises(DocumentParserError) as context:
                parse_document(path)
            self.assertEqual(context.exception.code, "unsupported_extension")

    def test_short_text_pdf_is_not_misclassified_as_scanned(self):
        import pymupdf

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "short.pdf"
            pdf = pymupdf.open()
            page = pdf.new_page()
            page.insert_text((72, 72), "short valid PDF")
            pdf.save(path)
            pdf.close()

            document = parse_document(path)
            self.assertEqual(document.parser, "pymupdf")
            self.assertEqual(document.metadata["page_count"], 1)
            self.assertIn("short valid PDF", document.text)

    def test_blank_pdf_requires_ocr(self):
        import pymupdf

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "blank.pdf"
            pdf = pymupdf.open()
            pdf.new_page()
            pdf.save(path)
            pdf.close()
            with self.assertRaises(DocumentParserError) as context:
                parse_document(path)
            self.assertEqual(context.exception.code, "scanned_pdf_requires_ocr")

    def test_docx_preserves_heading_paragraph_and_table_order(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "policy.docx"
            source = Document()
            source.add_heading("总则", level=1)
            source.add_paragraph("第一条内容")
            table = source.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "字段"
            table.cell(0, 1).text = "值"
            source.save(path)

            document = parse_document(path)
            self.assertEqual(document.parser, "python-docx")
            self.assertEqual([block.kind for block in document.blocks], ["heading", "paragraph", "table"])
            self.assertEqual(document.blocks[-1].heading_path, ["总则"])


if __name__ == "__main__":
    unittest.main()
